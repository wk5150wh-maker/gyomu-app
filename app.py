import streamlit as st
import io
import json
from google import genai
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

st.set_page_config(page_title="業務チェックシート自動生成", layout="centered")

# ============================
# 🔒 パスワード認証の仕組み
# ============================
def check_password():
    def password_entered():
        if st.session_state["password"] == "tatsuno2026": # 必要に応じてパスワードを変更
            st.session_state["password_correct"] = True
            del st.session_state["password"]
        else:
            st.session_state["password_correct"] = False

    if "password_correct" not in st.session_state or not st.session_state["password_correct"]:
        st.title("🔒 社内認証が必要です")
        st.text_input(
            "パスワードを入力してください", type="password", on_change=password_entered, key="password"
        )
        if "password_correct" in st.session_state and not st.session_state["password_correct"]:
            st.error("😕 パスワードが間違っています。")
        st.info("※このアプリは社内共有専用です。")
        return False
    return True

if not check_password():
    st.stop()

# ============================
# 本編アプリ（Secretsから自動でAPIキーを取得）
# ============================

st.title("📋 業務チェックシート自動生成アプリ")
st.write("正式企業名と電話番号を入力すると、AIが業務課題を分析し、洗練されたA4縦PowerPointチラシを生成します。")

# StreamlitのSecretsからAPIキーを自動取得
try:
    api_key = st.secrets["general"]["GEMINI_API_KEY"]
except Exception:
    st.error("⚠️ APIキーがStreamlitのSecretsに設定されていません。管理画面のSecrets設定を確認してください。")
    st.stop()

# 入力フォーム（APIキー入力欄をなくし、企業名と電話番号のみにする）
with st.form("input_form"):
    company_name = st.text_input("正式企業名", placeholder="例：株式会社タツノ")
    phone_number = st.text_input("電話番号", placeholder="例：0268-22-6001")
    
    submitted = st.form_submit_button("チラシ型チェックシートを生成する")

if submitted:
    if not company_name:
        st.error("企業名は必須です。")
    else:
        with st.spinner("Geminiが企業情報を分析し、チラシを作成中..."):
            try:
                # 1. Gemini クライアント初期化（Secretsのキーを使用）
                client = genai.Client(api_key=api_key)
                
                prompt = f"""
                企業名: {company_name}
                電話番号: {phone_number}
                
                上記の企業の公開情報を基に、日常業務で行われている作業（報告や案内を作る、資料や情報を読む・探す、メモや情報をまとめる、確認や相談の準備をする、など）を4つのカテゴリーに分けて洗い出し、必ず以下の純粋なJSON形式のみで出力してください。Markdownのコードブロック（```json ... ```）や前置き・解説は一切含めず、プレーンなJSON文字列だけを返してください。

                JSONフォーマット:
                {{
                  "categories": [
                    {{
                      "category_name": "報告や案内を作る",
                      "items": [
                        {{"item": "現場のメモや写真から、報告書をまとめ直すことがある", "memo": "社内メモ1"}},
                        {{"item": "日報や顧客向けの案内文・メールをゼロから書いている", "memo": "社内メモ2"}}
                      ]
                    }},
                    {{
                      "category_name": "資料や情報を読む・探す",
                      "items": [
                        {{"item": "長い仕様書や過去の資料から、必要な箇所を探すのに時間がかかる", "memo": "社内メモ3"}},
                        {{"item": "過去の似た案件や対応履歴を引っ張り出して確認することがある", "memo": "社内メモ4"}}
                      ]
                    }},
                    {{
                      "category_name": "メモや情報をまとめる",
                      "items": [
                        {{"item": "打合せのメモや問い合わせ内容を、分かりやすく整理し直している", "memo": "社内メモ5"}},
                        {{"item": "複数の資料や情報を一つに突き合わせて確認・比較している", "memo": "社内メモ6"}}
                      ]
                    }},
                    {{
                      "category_name": "確認や相談の準備をする",
                      "items": [
                        {{"item": "経験者や上長へ確認・相談する前に、状況や聞きたいことを整理している", "memo": "社内メモ7"}},
                        {{"item": "伝達漏れがないか、提出前のチェック項目を見比べて確認している", "memo": "社内メモ8"}}
                      ]
                    }}
                  ]
                }}
                ※必ずカテゴリーは4つ、各カテゴリの項目は正確に2つずつにしてください。
                """
                
                response = client.models.generate_content(
                    model='gemini-3.8-flash',
                    contents=prompt,
                    config={
                        'response_mime_type': 'application/json',
                    },
                )
                
                data = json.loads(response.text)
                
                # --- デザインテンプレートによるパワポ生成 ---
                FONT = "Yu Gothic"
                NAVY, INK, MUTED = "0F4C81", "1F2937", "6B7280"
                ORANGE, YELLOW, PALE_BLUE = "F28C28", "FFC15E", "BBD4F0"

                COMPANY_STR = f"{company_name} 様 御中"
                TITLE1 = "ちょっと待って、その業務の問題"
                TITLE2 = "改善できるかもしれません。"
                SUB_PRE, SUB_KEY, SUB_POST = (
                    "いつもの仕事の中にある",
                    "『読む・探す・書く・まとめる・確認する』",
                    "。当てはまるものに✓を付けてください。",
                )

                cat_styles = [
                    ("1E6B3A", "E8F3EC"),
                    ("0E8C8C", "E3F4F4"),
                    ("6C4BA6", "EEEAF7"),
                    ("8A5A2B", "F6EDE3")
                ]

                CATEGORIES = []
                internal_memos = []
                
                for idx, cat_data in enumerate(data.get("categories", [])):
                    c_color, c_tint = cat_styles[idx % len(cat_styles)]
                    c_name = cat_data.get("category_name", "カテゴリー")
                    items_list = [item_obj["item"] for item_obj in cat_data.get("items", [])[:2]]
                    
                    for item_obj in cat_data.get("items", []):
                        internal_memos.append((item_obj["item"], item_obj.get("memo", "-")))
                        
                    CATEGORIES.append((c_name, c_color, c_tint, items_list))

                TOP3_PRE, TOP3_KEY = "特に「減らせたら助かる」と思うものを", "3つ"
                FOOTNOTE = "※公開情報をもとにした訪問前の仮説です。実際の業務状況を確認するためのチェックシートです。"

                SLIDE_W, SLIDE_H = Emu(7562088), Emu(10689336)
                W = SLIDE_W / 914400
                M = 0.5
                CW = W - M * 2

                def rgb(hex_):
                    return RGBColor.from_string(hex_)

                def set_font(run, size, color, bold=False):
                    run.font.size = Pt(size)
                    run.font.bold = bold
                    run.font.color.rgb = rgb(color)
                    run.font.name = FONT
                    rPr = run._r.get_or_add_rPr()
                    for tag in ("a:ea", "a:cs"):
                        el = rPr.find(qn(tag))
                        if el is None:
                            el = etree.SubElement(rPr, qn(tag))
                        el.set("typeface", FONT)

                def add_text(slide, x, y, w, h, runs, size, color=INK, bold=False,
                             anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT, name=None):
                    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                    if name:
                        tb.name = name
                    tf = tb.text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = anchor
                    p = tf.paragraphs[0]
                    p.alignment = align
                    if isinstance(runs, str):
                        runs = [(runs, color, bold)]
                    for text, c, b in runs:
                        r = p.add_run()
                        r.text = text
                        set_font(r, size, c, b)
                    return tb

                def no_effect(shape):
                    for ref in shape._element.iter(qn("a:effectRef")):
                        ref.set("idx", "0")

                def add_shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0,
                              radius=None, name=None):
                    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
                    if name:
                        sp.name = name
                    sp.shadow.inherit = False
                    no_effect(sp)
                    if fill:
                        sp.fill.solid()
                        sp.fill.fore_color.rgb = rgb(fill)
                    else:
                        sp.fill.background()
                    if line:
                        sp.line.color.rgb = rgb(line)
                        sp.line.width = Pt(line_w)
                    else:
                        sp.line.fill.background()
                    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                        sp.adjustments[0] = radius / min(w, h)
                    return sp

                def badge(slide, x, y, d, color, glyph, name):
                    add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=color, name=name + "_bg")
                    add_text(slide, x, y, d, d, glyph, 14, "FFFFFF", True,
                             align=PP_ALIGN.CENTER, name=name)

                prs = Presentation()
                prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
                s = prs.slides.add_slide(prs.slide_layouts[6])

                # ヘッダー
                add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, M, 0.5, CW, 2.1, fill=NAVY,
                          radius=0.18, name="Hero")
                add_text(s, M + 0.4, 0.68, 4, 0.4, COMPANY_STR, 14, PALE_BLUE, True, name="Company")
                add_text(s, M + 0.4, 1.15, 5.3, 0.5, TITLE1, 22, "FFFFFF", True, name="Title1")
                add_text(s, M + 0.4, 1.65, 5.3, 0.75, TITLE2, 28, YELLOW, True, name="Title2")
                add_shape(s, MSO_SHAPE.OVAL, W - M - 1.75, 0.85, 1.35, 1.35,
                          fill="3A6B9A", line="FFFFFF", line_w=0.75, name="BadgeBg")
                add_text(s, W - M - 1.75, 0.85, 1.35, 1.35, "✓", 54, "FFFFFF", True,
                         align=PP_ALIGN.CENTER, name="BadgeCheck")

                # サブタイトル
                add_text(s, M, 2.78, CW, 0.55,
                         [(SUB_PRE, INK, False), (SUB_KEY, NAVY, True), (SUB_POST, INK, False)],
                         12, name="Subtitle")

                # 2x2 カテゴリー枠
                gx, gy, ch, grid_y = 0.25, 0.22, 2.5, 3.5
                cw = (CW - gx) / 2
                for i, (title, color, tint, items) in enumerate(CATEGORIES):
                    x = M + (i % 2) * (cw + gx)
                    y = grid_y + (i // 2) * (ch + gy)
                    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw, ch, fill=tint,
                              radius=0.14, name=f"Card{i+1}")
                    badge(s, x + 0.2, y + 0.18, 0.5, color, str(i + 1), f"CardNo{i+1}")
                    add_text(s, x + 0.82, y + 0.18, cw - 0.95, 0.5, title, 14, color, True,
                             name=f"CardTitle{i+1}")
                    for j, text in enumerate(items):
                        ry = y + 0.85 + j * 0.8
                        add_shape(s, MSO_SHAPE.RECTANGLE, x + 0.22, ry + 0.04, 0.2, 0.2,
                                  fill="FFFFFF", line=color, line_w=1.5, name=f"Check{i+1}_{j+1}")
                        add_text(s, x + 0.55, ry - 0.06, cw - 0.7, 0.7, text, 12, INK,
                                 anchor=MSO_ANCHOR.TOP, name=f"Item{i+1}_{j+1}")

                # 下部記入欄
                py, ph = grid_y + ch * 2 + gy + 0.25, 2.0
                add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, M, py, CW, ph, fill="FFFFFF",
                          line=ORANGE, line_w=2, radius=0.14, name="Top3Panel")
                badge(s, M + 0.25, py + 0.15, 0.4, ORANGE, "★", "Top3Icon")
                add_text(s, M + 0.8, py + 0.1, CW - 1.1, 0.5,
                         [(TOP3_PRE, INK, True), (TOP3_KEY, ORANGE, True)], 14, name="Top3Title")
                for k, num in enumerate("①②③"):
                    ry = py + 0.72 + k * 0.4
                    add_text(s, M + 0.3, ry, 0.4, 0.34, num, 16, ORANGE, True,
                             align=PP_ALIGN.CENTER, name=f"Num{k+1}")
                    ln = s.shapes.add_connector(1, Inches(M + 0.85), Inches(ry + 0.34),
                                                Inches(M + 0.85 + CW - 1.2), Inches(ry + 0.34))
                    ln.line.color.rgb = rgb("C9CED6")
                    ln.line.width = Pt(1)
                    ln.name = f"WriteLine{k+1}"
                    no_effect(ln)

                # 注意書き
                add_text(s, M, py + ph + 0.1, CW, 0.3, FOOTNOTE, 10, MUTED, name="Footnote")

                file_stream = io.BytesIO()
                prs.save(file_stream)
                file_stream.seek(0)
                
                st.success("デザインされたチラシ型チェックシートの作成が完了しました！")
                
                st.download_button(
                    label="📥 PowerPointファイルをダウンロード",
                    data=file_stream,
                    file_name=f"{company_name}_業務改善チェックシート.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )
                
                st.markdown("---")
                st.subheader("📝 社内担当者向けメモ（AI活用支援案）")
                for item, memo in internal_memos:
                    st.write(f"- **チェック項目**: {item}")
                    st.write(f"  👉 **社内メモ**: {memo}")

            except Exception as e:
                st.error(f"エラーが発生しました: {e}")